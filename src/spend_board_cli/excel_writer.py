from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

HEADER = ["Date", "Marchant", "Montant"]


@dataclass(frozen=True)
class TransactionRow:
    date: str
    merchant: str
    amount: float


def append_to_excel(xlsx_path: str | Path, rows: list[TransactionRow]) -> None:
    path = Path(xlsx_path)

    if path.exists():
        wb = load_workbook(path)
        ws: Worksheet = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(HEADER)

    for row in rows:
        ws.append([row.date, row.merchant, row.amount])

    wb.save(path)
    wb.close()
